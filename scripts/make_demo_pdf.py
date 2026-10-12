"""生成 PDF 版演示资料（验证 PDF 解析链路与「第 N 页」引用）。

做法：先用 python-docx 生成排版过的 Word 文档，再用捆绑的 LibreOffice Kit 转成 PDF，
落到 data/docs/samples/。这样 PDF 里有真实文本层（不是图片），能被 pypdfium2 正确解析。

用法：
    & 'C:\\Users\\19181\\PyCharmMiscProject\\.venv\\Scripts\\python.exe' scripts\\make_demo_pdf.py
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

TARGET_DIR = ROOT / "data" / "docs" / "samples"
TARGET_PDF = TARGET_DIR / "M1136打印机安装与维护指南.pdf"
WORK_DIR = ROOT / "data" / "test_tmp" / "pdf-build"

NODE = Path("E:/DeepSeek Harness/resources/runtime/primary-runtime/dependencies/node/bin/node.exe")
CLI = Path(
    "E:/DeepSeek Harness/resources/app.asar.unpacked/dsh/node_modules/"
    "@deepseek-ai/libreoffice-kit/lib/cli.js"
)

# 正文按「页」组织：每页一个列表，便于在引用里出现「第 N 页」
PAGES: list[dict[str, object]] = [
    {
        "title": "一、开箱与安装",
        "paragraphs": [
            "云打 M1136 是面向中小企业办公场景的黑白激光打印机，标配 150 页进纸盒，"
            "支持 USB 2.0 与 Wi-Fi 2.4G 连接，随机附带启动硒鼓（型号 CT-M1136）。",
            "安装步骤：",
            "1. 拆除机身全部橙色运输胶带，取出硒鼓仓内的保护纸；",
            "2. 打开前盖装入硒鼓，沿导轨推到底并听到卡扣声；",
            "3. 拉开纸盒放入 A4 纸，纸张数量不超过最大刻度线，并贴合两侧限位卡扣；",
            "4. 通电开机，等待屏幕显示就绪；",
            "5. 访问官网下载中心安装「M1136 系列驱动」，按向导完成连接。",
            "注意：本机型仅支持 2.4G 频段 Wi-Fi，无法连接 5G 频段。",
        ],
    },
    {
        "title": "二、日常维护与耗材",
        "paragraphs": [
            "硒鼓型号为 CT-M1136，理论打印量约 1600 页（按 5% 覆盖率计算），建议使用原装耗材。",
            "更换硒鼓步骤：关闭电源并拔掉电源线，等待 30 秒；打开前盖取出旧硒鼓；"
            "新硒鼓左右轻摇 5 至 6 次并撕掉密封胶条；装入后合上前盖；"
            "开机后在「菜单 → 维护 → 硒鼓复位」中复位计数器。",
            "每月清洁一次进纸轮与分离垫可显著减少卡纸；禁止使用酒精、天那水等溶剂擦拭机身。",
        ],
    },
    {
        "title": "三、卡纸处理（重点）",
        "paragraphs": [
            "本机型报错码 E07 表示卡纸。处理步骤如下：",
            "1. 关闭打印机电源，打开前盖与后盖，找到卡住的纸张；",
            "2. 双手沿走纸方向缓慢、均匀地抽出纸张，禁止反向硬拽，避免撕碎残留；",
            "3. 检查走纸通道内是否有纸屑残留，可用干燥无绒布清理；",
            "4. 确认纸张平整、未受潮卷边，且不超过纸盒最大刻度线；",
            "5. 合上前后盖后开机，若仍提示 E07，说明传感器或进纸轮异常，需返厂检修。",
            "若一天内卡纸超过三次，请检查进纸轮是否磨损发亮；磨损时需更换进纸轮组件。",
        ],
    },
    {
        "title": "四、规格与保修提示",
        "paragraphs": [
            "打印速度 20 页/分钟，分辨率 600×600 dpi，接口 USB 2.0 与 Wi-Fi 2.4G。",
            "整机保修两年，硒鼓等耗材保修三个月，均自购买之日起计算。",
            "人为损坏、私自拆机、进水受潮、使用非原装耗材导致的故障不在免费保修范围内。",
            "需要退换货时请保留原包装、全部附件与购机发票。",
        ],
    },
]


def build_docx(path: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    document = Document()

    # 中文字体需要显式指定 eastAsia，否则渲染可能回退成方框
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "Microsoft YaHei",
    )

    document.add_heading("云打 M1136 激光打印机 · 安装与维护指南", level=0)
    document.add_paragraph("文档版本 v1.0 ｜ 适用于云打 M1136 系列 ｜ 内部售后资料")

    for page in PAGES:
        document.add_page_break()
        document.add_heading(str(page["title"]), level=1)
        for line in page["paragraphs"]:  # type: ignore[union-attr]
            document.add_paragraph(str(line))

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def convert_to_pdf(source: Path) -> Path:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    if TARGET_PDF.exists():
        TARGET_PDF.unlink()

    # LibreOffice Kit 的临时目录需要可写，这里指到工作区里
    env = {"TMP": str(WORK_DIR), "TEMP": str(WORK_DIR)}
    result = subprocess.run(
        [str(NODE), str(CLI), "convert", "--input", str(source), "--output", str(TARGET_PDF)],
        capture_output=True,
        text=True,
        timeout=300,
        env={**__import__("os").environ, **env},
    )
    if result.returncode != 0 or not TARGET_PDF.exists():
        raise RuntimeError(
            f"LibreOffice 转换失败（exit {result.returncode}）：{result.stderr[-500:]}"
        )
    return TARGET_PDF


def main() -> int:
    docx_path = WORK_DIR / "M1136打印机安装与维护指南.docx"
    build_docx(docx_path)
    print(f"已生成 Word：{docx_path}（{docx_path.stat().st_size} 字节）")

    if not NODE.exists() or not CLI.exists():
        print("未找到捆绑的 LibreOffice Kit，跳过 PDF 转换")
        return 1

    pdf_path = convert_to_pdf(docx_path)
    print(f"已生成 PDF：{pdf_path}（{pdf_path.stat().st_size} 字节）")

    # 立即校验文本层与页数
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(pdf_path))
    print(f"页数：{len(pdf)}")
    for index in range(len(pdf)):
        page = pdf[index]
        textpage = page.get_textpage()
        text = textpage.get_text_range().strip()
        print(f"  第 {index + 1} 页：{len(text)} 字，开头「{text[:30].replace(chr(10), ' ')}」")
        textpage.close()
        page.close()
    pdf.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
